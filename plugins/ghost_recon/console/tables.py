"""Case tables as files (spec §11 "Tablas"): findings, evidence, timeline and criteria — the rows of their JSON
endpoints, with the same filters — as CSV (UTF-8 with BOM, so Excel shows the accents) or XLSX (openpyxl, with the
Ghost Recon metadata of the pack). Pure: rows in, bytes out."""

from __future__ import annotations

import csv
import io
import json
import re
import tempfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, Optional, Sequence, Tuple

from ..core import SIGNATURE, ids
from ..core.reports.xlsx import _clean, rewrite_app_xml

CSV_MEDIA = "text/csv; charset=utf-8"
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLSX_CREATOR = "Ghost Recon (www.ghostrecon.ai)"
# Text a spreadsheet would run as a formula when it opens the CSV (CSV injection); a plain signed number stays as is.
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_NUMBER_RE = re.compile(r"[+-]?\d+(?:[.,]\d+)*")


class XlsxUnavailable(RuntimeError):
    """openpyxl (a python_dependency of the plugin, not of Hermes) is not installed here."""


@dataclass(frozen=True)
class Column:
    header: str
    value: Callable[[Dict[str, Any]], Any]


@dataclass(frozen=True)
class Table:
    title: str  # Spanish: the XLSX sheet and title
    columns: Sequence[Column]


def _f(key: str) -> Callable[[Dict[str, Any]], Any]:
    return lambda row: row.get(key)


def _audit(key: str) -> Callable[[Dict[str, Any]], Any]:
    return lambda row: ids.short_audit(row[key]) if row.get(key) else ""


def _listed(key: str) -> Callable[[Dict[str, Any]], Any]:
    def value(row: Dict[str, Any]) -> str:
        items = row.get(key) or []
        items = items if isinstance(items, list) else [items]
        return "; ".join(i if isinstance(i, str) else json.dumps(i, ensure_ascii=False) for i in items)
    return value


TABLES: Dict[str, Table] = {
    "findings": Table("Hallazgos", (
        Column("ID", _f("id")), Column("Tipo", _f("kind")), Column("Título", _f("title")),
        Column("Descripción", _f("description")), Column("Importe", _f("amount")), Column("Moneda", _f("currency")),
        Column("Riesgo", _f("risk")), Column("Confianza", _f("confidence")), Column("Etiqueta", _f("label")),
        Column("Estado", _f("status")), Column("Contraparte", _f("counterparty")), Column("Entidad", _f("entity")),
        Column("Fecha", _f("date")), Column("Categoría", _f("category")), Column("Quién aporta", _f("owner")),
        Column("Siguiente evidencia", _f("next_evidence")), Column("Evidencia", _listed("evidence_refs")),
        Column("Auditoría", _audit("audit_id")))),
    "evidence": Table("Evidencia", (
        Column("Ruta", _f("path")), Column("Archivo", _f("filename")), Column("Extensión", _f("ext")),
        Column("Tamaño (bytes)", _f("size")), Column("SHA-256", _f("sha256")), Column("MD5", _f("md5")),
        Column("Estado", _f("status")), Column("Bloque", _f("block")), Column("Tipo de documento", _f("doc_type")),
        Column("Fecha del documento", _f("doc_date")), Column("Entidad", _f("entity")),
        Column("Revisión", _f("review_status")), Column("Primera auditoría", _audit("first_audit_id")),
        Column("Dentro de un ZIP", lambda row: "sí" if row.get("zip_member") else "no"))),
    "timeline": Table("Cronología", (
        Column("Fecha", _f("ts")), Column("Evento", _f("event_type")), Column("Auditoría", _audit("audit_id")),
        Column("Actor", _f("actor")), Column("Descripción", _f("description")))),
    "criteria": Table("Criterios", (
        Column("ID", _f("id")), Column("Fecha", _f("date")), Column("Autor", _f("author")),
        Column("Texto (literal)", _f("text")), Column("Estado", _f("status")), Column("Auditoría", _audit("audit_id")))),
}


def csv_cell(value: Any) -> Any:
    """A CSV value: empty for None, JSON for structures, and text that a spreadsheet would read as a formula behind a
    leading apostrophe (OWASP CSV injection), unless it is just a signed number."""
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value.startswith(_FORMULA_START) and not _NUMBER_RE.fullmatch(value):
        return "'" + value
    return value


def to_csv(table: Table, rows: Iterable[Dict[str, Any]], case: Dict[str, Any]) -> bytes:
    """RFC 4180 (comma, CRLF, quoted fields) in UTF-8 with a BOM."""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow([c.header for c in table.columns])
    writer.writerows([csv_cell(c.value(row)) for c in table.columns] for row in rows)
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def load_openpyxl():
    """Imported only when an XLSX is asked for: without it the console still serves everything else."""
    try:
        import openpyxl
    except ImportError as exc:
        raise XlsxUnavailable("Falta openpyxl, dependencia del plugin Ghost Recon: la exportación a XLSX no está "
                              "disponible en esta instalación (CSV sí lo está).") from exc
    return openpyxl


def to_xlsx(table: Table, rows: Iterable[Dict[str, Any]], case: Dict[str, Any]) -> bytes:
    """One sheet, bold frozen header; Ghost Recon properties and docProps/app.xml, as the pack's workbook."""
    openpyxl = load_openpyxl()
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = table.title
    ws.append([c.header for c in table.columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for r, row in enumerate(rows, start=2):
        for col, column in enumerate(table.columns, start=1):
            value = column.value(row)
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            value = _clean(value)  # OCR and chat exports carry control characters (one cleaner, shared with the pack)
            cell = ws.cell(row=r, column=col, value=value)
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"  # data that looks like a formula is text, never a formula (as in the pack)
    props = wb.properties
    props.creator = XLSX_CREATOR
    props.lastModifiedBy = "Ghost Recon"
    props.title = f"Ghost Recon — {table.title} — {case['name']}"
    props.subject = "Auditoría financiera forense"
    props.description = SIGNATURE
    props.keywords = f"Ghost Recon, forensic audit, {table.title}"
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "table.xlsx"
        wb.save(str(path))
        rewrite_app_xml(path)
        return path.read_bytes()


WRITERS: Dict[str, Tuple[Callable[[Table, Iterable[Dict[str, Any]], Dict[str, Any]], bytes], str]] = {
    "csv": (to_csv, CSV_MEDIA), "xlsx": (to_xlsx, XLSX_MEDIA)}


def file_name(case: Dict[str, Any], table: str, fmt: str, now: Optional[datetime] = None) -> str:
    """``GhostRecon_<slug>_<table>_<YYYYMMDD-HHMM>.<fmt>`` (UTC)."""
    stamp = (now or datetime.now(timezone.utc)).strftime("%Y%m%d-%H%M")
    return f"GhostRecon_{case['slug']}_{table}_{stamp}.{fmt}"
