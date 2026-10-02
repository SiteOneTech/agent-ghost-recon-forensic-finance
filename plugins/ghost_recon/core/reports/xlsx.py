"""Evidence workbook (XLSX) with openpyxl: README, DOCUMENT_INDEX, FINDINGS registers, CRITERIA, AUDIT_TRAIL,
VERSION_EFFECT, RESEARCH, DASHBOARD — only the sheets with content. Properties and docProps/app.xml say Ghost Recon.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from .. import ENGINE, SIGNATURE, ids

NAVY = "1A2A4D"
BAND = "F0F2F6"
AMBER = "FFF2CC"
RED = "F8D7DA"
GREEN = "D4EDDA"
MONEY = '#,##0.00'


def _style(ws, header_row: int, ncols: int, widths: Optional[List[int]] = None):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    for c in range(1, ncols + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _title(ws, text: str, subtitle: str):
    from openpyxl.styles import Font
    ws["A1"] = text; ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = subtitle; ws["A2"].font = Font(italic=True, size=9, color="666666")


def _clean(v: Any) -> Any:
    """openpyxl refuses control characters (OCR / chat exports carry them); strip them from strings."""
    if isinstance(v, str):
        try:
            from openpyxl.utils.cell import ILLEGAL_CHARACTERS_RE
            return ILLEGAL_CHARACTERS_RE.sub("", v)
        except Exception:
            return v
    return v


def _put(ws, row: int, values: List[Any], money_cols: Optional[List[int]] = None):
    from openpyxl.styles import Alignment
    for c, v in enumerate(values, 1):
        v = _clean(v)
        cell = ws.cell(row=row, column=c, value=v)
        if isinstance(v, str) and v.startswith("="):  # data that looks like a formula is text, never a formula
            cell.data_type = "s"
        if money_cols and c in money_cols and isinstance(v, (int, float)):
            cell.number_format = MONEY
        cell.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)


def _band_by_status(ws, row: int, status: str, ncols: int):
    from openpyxl.styles import PatternFill
    color = {"open": AMBER, "closed": GREEN, "downgraded": BAND, "UNRESOLVED": RED, "CONFIRMED": GREEN}.get(status)
    if color:
        for c in range(1, ncols + 1):
            ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor=color)


def _footer(wb, text: str):
    for ws in wb.worksheets:
        ws.oddFooter.left.text = text
        ws.oddFooter.right.text = "Página &P de &N"


def build_workbook(case: Dict[str, Any], audit: Dict[str, Any], *, dest: Path, version: str, model: Dict[str, Any],
                   findings: List[Dict[str, Any]], criteria: List[Dict[str, Any]], evidence: List[Dict[str, Any]],
                   research: List[Dict[str, Any]], validation: Optional[Dict[str, Any]] = None,
                   reports: Optional[List[Dict[str, Any]]] = None) -> Path:
    from openpyxl import Workbook
    wb = Workbook()
    cur = case.get("base_currency", "USD")
    short = ids.short_audit(audit["id"])
    sub = f"{case['name']} · {short} ({audit['kind']}) · {version} · {datetime.now().date().isoformat()} · moneda base {cur}"

    # 00_README
    ws = wb.active; ws.title = "00_README"
    _title(ws, f"Ghost Recon — Workbook de evidencia · {case['name']}", sub)
    rows = [("Caso", case["id"]), ("Carpeta de evidencia", case["root_path"]), ("Auditoría", audit["id"]),
            ("Tipo", audit["kind"]), ("Carpeta de la auditoría", audit["folder"]), ("Versión", version),
            ("Moneda base", cur), ("Contexto", audit.get("context_md") or ""),
            ("Archivos de evidencia", len(evidence)), ("Excepciones", sum(1 for f in findings if f["kind"] == "exception")),
            ("Excepciones abiertas", sum(1 for f in findings if f["kind"] == "exception" and f["status"] == "open")),
            ("Anomalías", sum(1 for f in findings if f["kind"] == "anomaly")),
            ("Preguntas abiertas", sum(1 for f in findings if f["kind"] == "question" and f["status"] == "open")),
            ("Criterios registrados", len(criteria))]
    r = 4
    _put(ws, r, ["Campo", "Valor"]); _style(ws, r, 2, [34, 90]); r += 1
    for k, v in rows:
        _put(ws, r, [k, v]); r += 1
    r += 1
    _put(ws, r, ["Hoja", "Contenido"]); _style(ws, r, 2); r += 1
    index_row = r
    r += 12
    if validation and validation.get("rounds"):
        _put(ws, r, ["Validación independiente", ""]); r += 1
        for rd in validation["rounds"]:
            _put(ws, r, [f"Auditor {rd.get('role')}", f"{rd.get('status')} — {rd.get('summary', '')}"]); r += 1
    r += 1
    _put(ws, r, ["Firma", SIGNATURE])

    # 02_DOCUMENT_INDEX
    ws = wb.create_sheet("02_DOCUMENT_INDEX")
    _title(ws, "Índice de documentos (un archivo por fila; miembros de ZIP incluidos)", sub)
    hdr = ["Ruta", "Nombre", "Ext", "Tamaño", "Fecha mod.", "SHA-256", "MD5", "ZIP", "Bloque", "Estado", "Primera auditoría", "Tipo doc", "Revisión", "Confianza", "Notas"]
    _put(ws, 4, hdr); _style(ws, 4, len(hdr), [50, 28, 6, 10, 20, 66, 34, 5, 14, 12, 14, 16, 12, 12, 30])
    for i, e in enumerate(evidence, 5):
        _put(ws, i, [e["path"], e["filename"], e["ext"], e["size"], e.get("mtime"), e["sha256"], e.get("md5"), "sí" if e.get("zip_member") else "",
                     e.get("block"), e.get("status"), ids.short_audit(e["first_audit_id"]) if e.get("first_audit_id") else "",
                     e.get("doc_type"), e.get("review_status"), e.get("confidence"), e.get("notes")])

    # Registers
    kinds = [("30_EXCEPTIONS", "exception", "Registro de excepciones"), ("29_ANOMALIES", "anomaly", "Registro de anomalías"),
             ("33_QUESTIONS", "question", "Preguntas para quien tiene la evidencia"), ("34_FINDINGS", "finding", "Hallazgos")]
    for sheet, kind, label in kinds:
        rows_k = [f for f in findings if f["kind"] == kind]
        if not rows_k:
            continue
        ws = wb.create_sheet(sheet)
        _title(ws, label, sub)
        hdr = ["ID", "Título", "Descripción", "Importe", "Moneda", "Entidad", "Contraparte", "Fecha", "Categoría", "Riesgo", "Confianza",
               "Etiqueta", "Estado", "Evidencia (refs)", "Evidencia que lo resolvería", "Quién la aporta", "Última auditoría", "Historia"]
        _put(ws, 4, hdr); _style(ws, 4, len(hdr), [9, 36, 50, 14, 7, 16, 18, 11, 14, 9, 14, 12, 10, 24, 34, 16, 10, 40])
        for i, f in enumerate(rows_k, 5):
            hist = "; ".join(f"{ids.short_audit(h['audit_id']) if h.get('audit_id') else '-'} {h.get('ts', '')[:10]}: "
                             f"{h.get('change') if isinstance(h.get('change'), str) else ', '.join(h.get('change', {}).keys())}" for h in (f.get("history") or []))
            _put(ws, i, [f["id"], f["title"], f.get("description"), f.get("amount"), f.get("currency") or cur, f.get("entity"), f.get("counterparty"),
                         f.get("date"), f.get("category"), f["risk"], f["confidence"], f["label"], f["status"],
                         ", ".join(map(str, f.get("evidence_refs") or [])), f.get("next_evidence"), f.get("owner"),
                         ids.short_audit(f["audit_id"]) if f.get("audit_id") else "", hist], money_cols=[4])
            _band_by_status(ws, i, f["status"], len(hdr))
        n = len(rows_k)
        ws.cell(row=n + 6, column=1, value="Totales (calculados)")
        ws.cell(row=n + 6, column=2, value=f'=COUNTA(A5:A{n + 4})')
        ws.cell(row=n + 6, column=3, value=f'=COUNTIF(M5:M{n + 4},"open")')
        ws.cell(row=n + 6, column=4, value=f'=SUM(D5:D{n + 4})').number_format = MONEY

    if criteria:
        ws = wb.create_sheet("CRIT_CRITERIA")
        _title(ws, "Criterios registrados (instrucciones/declaraciones de las partes — no conclusiones de la auditoría)", sub)
        hdr = ["ID", "Fecha", "Autor", "Texto literal", "Estado", "Auditoría"]
        _put(ws, 4, hdr); _style(ws, 4, len(hdr), [9, 11, 18, 90, 20, 10])
        for i, c in enumerate(criteria, 5):
            _put(ws, i, [c["id"], c.get("date"), c["author"], c["text"], c["status"], ids.short_audit(c["audit_id"]) if c.get("audit_id") else ""])

    trail = model.get("audit_trail") or []
    if trail:
        ws = wb.create_sheet("31_AUDIT_TRAIL")
        _title(ws, "Pista de auditoría: cifra cabecera → hoja → documentos → confianza → método", sub)
        hdr = ["Cifra", "Valor", "Hoja / sección", "Documentos", "Confianza", "Método / nota"]
        _put(ws, 4, hdr); _style(ws, 4, len(hdr), [34, 16, 22, 50, 14, 50])
        for i, t in enumerate(trail, 5):
            _put(ws, i, [t.get("figure"), t.get("value"), t.get("sheet"), ", ".join(map(str, t.get("documents") or [])), t.get("confidence"), t.get("method")], money_cols=[2])

    ve = model.get("version_effect") or []
    if ve:
        ws = wb.create_sheet(f"3x_VERSION_EFFECT_{version}")
        _title(ws, "Efecto de versión: partida · anterior · nuevo · diferencia · documento causante", sub)
        hdr = ["Partida", "Anterior", "Nuevo", "Diferencia", "Documento causante", "Hoja", "Nota"]
        _put(ws, 4, hdr); _style(ws, 4, len(hdr), [36, 16, 16, 16, 40, 16, 40])
        for i, v in enumerate(ve, 5):
            _put(ws, i, [v.get("item"), v.get("previous"), v.get("new"), v.get("difference"), v.get("document"), v.get("sheet"), v.get("note")], money_cols=[2, 3, 4])

    if research:
        ws = wb.create_sheet("35_RESEARCH")
        _title(ws, "Investigación externa registrada (los resultados web son INFERENCE salvo documento oficial)", sub)
        hdr = ["#", "Acción", "Consulta", "URL", "Título", "Extracto", "Relevancia", "Archivo"]
        _put(ws, 4, hdr); _style(ws, 4, len(hdr), [5, 10, 40, 50, 40, 60, 10, 40])
        for i, n in enumerate(research, 5):
            _put(ws, i, [n["id"], n["action"], n["query"], n.get("url"), n.get("title"), n.get("snippet"), n.get("relevance"), n.get("saved_path")])

    # DASHBOARD
    ws = wb.create_sheet("32_DASHBOARD")
    _title(ws, "Dashboard (cifras del modelo; conteos calculados desde los registros)", sub)
    _put(ws, 4, ["Indicador", "Valor", "Fuente"]); _style(ws, 4, 3, [44, 20, 40])
    r = 5
    for k, v in (model.get("kpis") or {}).items():
        _put(ws, r, [k, v, "model.json/kpis"], money_cols=[2]); r += 1
    counts = [("Excepciones (total)", "30_EXCEPTIONS", "exception"), ("Anomalías (total)", "29_ANOMALIES", "anomaly"),
              ("Preguntas (total)", "33_QUESTIONS", "question")]
    for label, sheet, kind in counts:
        n = sum(1 for f in findings if f["kind"] == kind)
        if n:
            ws.cell(row=r, column=1, value=label); ws.cell(row=r, column=2, value=f"=COUNTA('{sheet}'!A5:A{n + 4})"); ws.cell(row=r, column=3, value=sheet); r += 1
            ws.cell(row=r, column=1, value=label.replace("(total)", "abiertas")); ws.cell(row=r, column=2, value=f"=COUNTIF('{sheet}'!M5:M{n + 4},\"open\")"); ws.cell(row=r, column=3, value=sheet); r += 1
    _put(ws, r, ["Archivos de evidencia", len(evidence), "02_DOCUMENT_INDEX"]); r += 1
    _put(ws, r, ["Números que no deben mezclarse", "utilidad ≠ caja ≠ balance ≠ posición de socios", "método §1"])

    # sheet index back in README
    ws0 = wb["00_README"]
    for i, name in enumerate(wb.sheetnames):
        _put(ws0, index_row + i, [name, ""])
    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 100
    _footer(wb, f"{SIGNATURE} · {short} {version}")
    wb.properties.creator = "Ghost Recon (www.ghostrecon.ai)"
    wb.properties.lastModifiedBy = "Ghost Recon"
    wb.properties.title = f"Ghost Recon — Workbook de evidencia — {case['name']} — {short} {version}"
    wb.properties.subject = "Auditoría financiera forense"
    wb.properties.description = SIGNATURE
    wb.properties.keywords = "Ghost Recon, forensic audit, evidence workbook"
    dest = Path(dest); dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(dest))
    rewrite_app_xml(dest)
    return dest


def build_chronology_workbook(chron: Dict[str, Any], dest: Path) -> Path:
    from openpyxl import Workbook
    wb = Workbook()
    case = chron["case"]
    ws = wb.active; ws.title = "00_AUDITS"
    _title(ws, f"Cronología del caso {case['name']}", case["id"])
    hdr = ["Auditoría", "Tipo", "Estado", "Inicio", "Sello", "Evidencia nueva", "Carpeta"]
    _put(ws, 4, hdr); _style(ws, 4, len(hdr), [12, 10, 12, 20, 20, 14, 70])
    for i, a in enumerate(chron["audits"], 5):
        s = a.get("summary") or {}
        _put(ws, i, [ids.short_audit(a["id"]), a["kind"], a["status"], a["started_at"], a.get("sealed_at"), s.get("evidence_new"), a["folder"]])
    ws = wb.create_sheet("01_TIMELINE")
    _title(ws, "Línea de tiempo", case["id"])
    hdr = ["Fecha", "Tipo", "Auditoría", "Evento"]
    _put(ws, 4, hdr); _style(ws, 4, len(hdr), [20, 20, 10, 110])
    for i, e in enumerate(chron["entries"], 5):
        _put(ws, i, [e.get("ts"), e["type"], ids.short_audit(e["audit_id"]) if e.get("audit_id") else "", e["text"]])
    ws = wb.create_sheet("02_FINDINGS")
    _title(ws, "Evolución de hallazgos", case["id"])
    hdr = ["ID", "Tipo", "Título", "Estado", "Riesgo", "Confianza", "Importe", "Historia"]
    _put(ws, 4, hdr); _style(ws, 4, len(hdr), [9, 10, 40, 10, 9, 14, 14, 80])
    for i, f in enumerate(chron["findings"], 5):
        hist = "; ".join(f"{ids.short_audit(h['audit_id']) if h.get('audit_id') else '-'}: {h.get('change') if isinstance(h.get('change'), str) else ', '.join(h.get('change', {}).keys())}" for h in (f.get("history") or []))
        _put(ws, i, [f["id"], f["kind"], f["title"], f["status"], f["risk"], f["confidence"], f.get("amount"), hist], money_cols=[7])
    if chron["criteria"]:
        ws = wb.create_sheet("03_CRITERIA")
        _title(ws, "Criterios registrados", case["id"])
        hdr = ["ID", "Fecha", "Autor", "Texto", "Estado"]
        _put(ws, 4, hdr); _style(ws, 4, len(hdr), [9, 11, 18, 90, 20])
        for i, c in enumerate(chron["criteria"], 5):
            _put(ws, i, [c["id"], c.get("date"), c["author"], c["text"], c["status"]])
    _footer(wb, SIGNATURE)
    wb.properties.creator = "Ghost Recon (www.ghostrecon.ai)"; wb.properties.lastModifiedBy = "Ghost Recon"
    wb.properties.title = f"Ghost Recon — Cronología — {case['name']}"; wb.properties.description = SIGNATURE
    dest = Path(dest); dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(dest)); rewrite_app_xml(dest)
    return dest


def rewrite_app_xml(path: Path) -> None:
    """openpyxl stamps docProps/app.xml with its own Application name; rewrite it in place."""
    path = Path(path)
    src = path.read_bytes()
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                text = data.decode("utf-8", errors="ignore")
                text = re.sub(r"<Application>.*?</Application>", f"<Application>{ENGINE}</Application>", text)
                text = re.sub(r"<AppVersion>.*?</AppVersion>", "<AppVersion>1.0</AppVersion>", text)
                if "<Application>" not in text:
                    text = text.replace("</Properties>", f"<Application>{ENGINE}</Application></Properties>")
                data = text.encode("utf-8")
            zout.writestr(item, data)
    path.write_bytes(buf.getvalue())
